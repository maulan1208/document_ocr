from io import BytesIO
import unittest

from docx import Document

from app.services.document_processor import (
    DocumentProcessingError,
    process_document,
)


class DocumentProcessorTests(unittest.TestCase):
    def test_processes_utf8_txt_with_bom(self):
        result = process_document(
            "\ufeffNội dung tiếng Việt".encode("utf-8"),
            "txt",
        )

        self.assertEqual(result, ("Nội dung tiếng Việt", 1, "text"))

    def test_processes_utf16_txt_with_bom(self):
        result = process_document(
            "Nội dung tiếng Việt".encode("utf-16"),
            "txt",
        )

        self.assertEqual(result, ("Nội dung tiếng Việt", 1, "text"))

    def test_rejects_txt_with_unsupported_encoding(self):
        with self.assertRaises(DocumentProcessingError):
            process_document(b"\xff invalid utf-8", "txt")

    def test_processes_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Nội dung tiếng Việt")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Tên"
        table.cell(0, 1).text = "Giá trị"
        file_bytes = BytesIO()
        document.save(file_bytes)

        result = process_document(file_bytes.getvalue(), "docx")

        self.assertEqual(
            result,
            ("Nội dung tiếng Việt\nTên | Giá trị", 1, "text"),
        )

    def test_rejects_invalid_docx(self):
        with self.assertRaises(DocumentProcessingError):
            process_document(b"not a DOCX file", "docx")


if __name__ == "__main__":
    unittest.main()
